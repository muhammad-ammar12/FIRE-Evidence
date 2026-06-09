from docx import Document



def write_into_word_file(pico_evidence,json_stats,fileName='output_word_file'):

    doc=Document()

    doc.add_heading('PICO variables', level=1)
    doc.add_paragraph(pico_evidence)


    doc.add_heading('Main results (Stats)', level=1)
    doc.add_paragraph(json_stats)

    
    doc.save(fileName)

    return print(fileName + " has been saved in a word file.")