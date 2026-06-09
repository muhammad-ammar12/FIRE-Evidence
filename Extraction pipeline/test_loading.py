import fitz  # PyMuPDF
import pdfplumber
import pytesseract
from pdf2image import convert_from_path
from PIL import Image
from docx import Document
from docx.shared import Pt
from docx.oxml.ns import qn
import io
import os
import re 

def clean_text(text):
    if not isinstance(text, str):
        return ""
    text = text.replace('\x00', '')  # Remove NULL bytes
    text = re.sub(r'[\x01-\x08\x0B-\x0C\x0E-\x1F]', '', text)  # Remove control characters
    return text.strip()

def is_scanned_pdf(pdf_path, sample_pages=2):
    """Check if a PDF is scanned (image-based) by looking for no extractable text."""
    doc = fitz.open(pdf_path)
    for page_num in range(min(sample_pages, len(doc))):
        page = doc.load_page(page_num)
        if page.get_text("text").strip():
            return False  # Found text
    return True

def extract_text_with_layout(pdf_path):
    doc = fitz.open(pdf_path)
    text_data = []
    for page_num in range(len(doc)):
        page = doc.load_page(page_num)
        blocks = page.get_text("blocks")
        blocks.sort(key=lambda b: (b[1], b[0]))  # sort by vertical then horizontal
        for b in blocks:
            text_data.append({'type': 'text', 'page': page_num + 1, 'content': b[4].strip()})
    return text_data

def extract_tables(pdf_path):
    tables = []
    with pdfplumber.open(pdf_path) as pdf:
        for page_num, page in enumerate(pdf.pages):
            page_tables = page.extract_tables()
            for table in page_tables:
                tables.append({'type': 'table', 'page': page_num + 1, 'content': table})
    return tables

def extract_ocr_text(pdf_path):
    images = convert_from_path(pdf_path)
    ocr_data = []
    for i, image in enumerate(images):
        text = pytesseract.image_to_string(image)
        ocr_data.append({'type': 'text', 'page': i + 1, 'content': text.strip()})
    return ocr_data

def extract_pdf_structured(pdf_path):
    print(f"Analyzing '{pdf_path}' ...")
    scanned = is_scanned_pdf(pdf_path)
    print(f"PDF is {'scanned (image-based)' if scanned else 'digitally generated'}.")

    text_data = extract_ocr_text(pdf_path) if scanned else extract_text_with_layout(pdf_path)
    table_data = extract_tables(pdf_path)

    # Combine by page number
    structured_output = {}
    for item in text_data + table_data:
        page = item['page']
        structured_output.setdefault(page, []).append(item)

    return structured_output

def print_structured_output(structured_data):
    for page in sorted(structured_data.keys()):
        print(f"\n=== Page {page} ===")
        for item in structured_data[page]:
            if item['type'] == 'text':
                print("\n[Text Block]\n", item['content'])
            elif item['type'] == 'table':
                print("\n[Table]")
                for row in item['content']:
                    print(" | ".join(cell if cell else "" for cell in row))



def write_to_word(structured_data, output_path="saved_output.docx"):
    doc = Document()
    style = doc.styles['Normal']
    font = style.font
    font.name = 'Calibri'
    font.size = Pt(11)

    for page in sorted(structured_data.keys()):
        doc.add_heading(f'Page {page}', level=1)

        for item in structured_data[page]:
            if item['type'] == 'text':
                # 🔧 CHANGED: Sanitize text before adding
                doc.add_paragraph(clean_text(item['content']))
            elif item['type'] == 'table':
                table_data = item['content']
                if not table_data:
                    continue
                rows = len(table_data)
                cols = max(len(r) for r in table_data)
                table = doc.add_table(rows=rows, cols=cols)
                table.style = 'Table Grid'

                for i, row in enumerate(table_data):
                    for j, cell in enumerate(row):
                        # 🔧 CHANGED: Clean table cell text
                        table.cell(i, j).text = clean_text(cell) if cell else ""

        doc.add_page_break()

    doc.save(output_path)
    print(f"✅ Word document saved as '{output_path}'")

