import latest_LLM as lm
from docx import Document



def run_standardization(pico_text:str):

    graph_entities=lm.up_Format_evidence(pico_text)
    #graph_entities=format.ensure_ids_and_numbers(model=graph_entities,pico_text=pico_text)
    print(graph_entities)
    document = Document()
    document.add_heading("FHIR Compliant Standard Knowledge (Computable)", level=1)
    document.add_paragraph(str(graph_entities))

    # Save file
    document.save("batur-2025_computable.docx")







